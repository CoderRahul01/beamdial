"""Build results/BeamDial-results-v0.1.xlsx from graded runs. All summary numbers are live formulas."""
import collections
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
L = ["low", "medium", "high", "xhigh", "max"]
G = {"gsm8k": "GSM8K", "math_l1": "MATH", "math_l2": "MATH", "math_l3": "MATH", "math_l4": "MATH", "math_l5": "MATH",
     "amc23": "Competition", "aime24": "Competition", "olympiad": "Competition", "humaneval": "HumanEval", "mbpp": "MBPP"}
tasks = {t["id"]: t for t in map(json.loads, open(ROOT / "data" / "tasks.jsonl"))}
D = collections.defaultdict(dict)
for r in map(json.loads, open(ROOT / "data" / "graded.jsonl")):
    D[r["task_id"]][r["effort"]] = r
ids = sorted(D)
flags = json.loads((ROOT / "data" / "flagged_tasks.json").read_text())["tasks"]
summary = json.loads((ROOT / "results" / "summary.json").read_text())

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F2937")
B = Font(name=F, bold=True)
N = Font(name=F)
INP = Font(name=F, color="0000FF")
YEL = PatternFill("solid", fgColor="FFFF00")
NOTE = Font(name=F, italic=True, color="666666", size=9)
thin = Side(style="thin", color="D0D0D0")
GRID = Border(bottom=thin)


def header(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font, cell.fill = H, HF
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


wb = Workbook()

# ---------- Per task (raw data) ----------
pt = wb.active
pt.title = "Per task"
cols = ["Task", "Source", "Group", "Family", "Flagged"] + [f"Correct {l}" for l in L] + [f"Tokens {l}" for l in L] + [f"Reasoning {l}" for l in L]
header(pt, 1, cols, [26, 12, 13, 9, 9] + [10] * 15)
for i, t in enumerate(ids, 2):
    src = tasks[t]["source"]
    row = [t, src, G[src], tasks[t]["family"], "yes" if t in flags else "no"]
    row += [int(D[t][l]["correct"]) for l in L]
    row += [D[t][l]["usage"]["completion_tokens"] for l in L]
    row += [D[t][l]["usage"].get("reasoning_tokens") or 0 for l in L]
    for j, v in enumerate(row, 1):
        c = pt.cell(row=i, column=j, value=v)
        c.font = INP if j > 5 else N
pt.freeze_panes = "B2"
NT = len(ids) + 1  # last data row
pt.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{NT}"
corr_col = {l: get_column_letter(6 + k) for k, l in enumerate(L)}
tok_col = {l: get_column_letter(11 + k) for k, l in enumerate(L)}
rsn_col = {l: get_column_letter(16 + k) for k, l in enumerate(L)}

# ---------- Summary ----------
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "BeamDial v0.1: Beam reasoning levels on 410 graded tasks"
sm["A1"].font = Font(name=F, bold=True, size=14)
sm["A2"] = ("Model: Beam 1.0 (501B A23B) via Hugging Face, provider Baseten. 1 sample per task per level, temperature 0.7. "
            "Every number below is a formula over the 'Per task' sheet.")
sm["A2"].font = NOTE
sm["A4"], sm["A5"], sm["A6"] = "Inputs", "Input price (USD per 1M tokens)", "Output price (USD per 1M tokens)"
sm["A4"].font = B
sm["B5"], sm["B6"] = 0.40, 1.20
for c in ("B5", "B6"):
    sm[c].font, sm[c].fill, sm[c].number_format = INP, YEL, "$0.00"
sm["C5"] = "Change these to your provider's real price. Assumed list price, check your Hugging Face billing page."
sm["C5"].font = NOTE
sm["A7"] = "Mean prompt tokens per task"
sm["B7"] = round(sum(D[t]["low"]["usage"]["prompt_tokens"] for t in ids) / len(ids))
sm["B7"].font = INP
sm["C7"] = "Measured from the runs (same prompt at every level)."
sm["C7"].font = NOTE

r0 = 9
header(sm, r0, ["Level", "Accuracy", "95% CI low", "95% CI high", "Accuracy (clean tasks)", "Mean completion tokens",
                "Mean reasoning tokens", "Tokens vs max", "USD per 1,000 tasks"], [34, 12, 12, 12, 14, 14, 14, 12, 14])
for k, l in enumerate(L):
    r = r0 + 1 + k
    cc, tc, rc = corr_col[l], tok_col[l], rsn_col[l]
    sm.cell(row=r, column=1, value=l).font = B
    sm.cell(row=r, column=2, value=f"=AVERAGE('Per task'!{cc}2:{cc}{NT})")
    n = f"COUNT('Per task'!{cc}2:{cc}{NT})"
    p = f"B{r}"
    z = "1.96"
    centre = f"(({p}+{z}^2/(2*{n}))/(1+{z}^2/{n}))"
    half = f"({z}*SQRT({p}*(1-{p})/{n}+{z}^2/(4*{n}^2))/(1+{z}^2/{n}))"
    sm.cell(row=r, column=3, value=f"={centre}-{half}")
    sm.cell(row=r, column=4, value=f"={centre}+{half}")
    sm.cell(row=r, column=5, value=f"=AVERAGEIFS('Per task'!{cc}2:{cc}{NT},'Per task'!$E$2:$E${NT},\"no\")")
    sm.cell(row=r, column=6, value=f"=AVERAGE('Per task'!{tc}2:{tc}{NT})")
    sm.cell(row=r, column=7, value=f"=AVERAGE('Per task'!{rc}2:{rc}{NT})")
    sm.cell(row=r, column=8, value=f"=F{r}/$F${r0 + 5}")
    sm.cell(row=r, column=9, value=f"=1000*($B$7*$B$5+F{r}*$B$6)/1000000")
    for j, fmt in zip(range(2, 10), ["0.0%", "0.0%", "0.0%", "0.0%", "#,##0", "#,##0", "0%", "$0.00"]):
        sm.cell(row=r, column=j).number_format = fmt
        sm.cell(row=r, column=j).font = N
re_ = r0 + 7
sm.cell(row=re_, column=1, value="Head to head, low vs max").font = B
sm.cell(row=re_ + 1, column=1, value="low wrong, max right (tasks)")
sm.cell(row=re_ + 1, column=2, value=f"=SUMPRODUCT(('Per task'!{corr_col['low']}2:{corr_col['low']}{NT}=0)*('Per task'!{corr_col['max']}2:{corr_col['max']}{NT}=1))")
sm.cell(row=re_ + 2, column=1, value="low right, max wrong (tasks)")
sm.cell(row=re_ + 2, column=2, value=f"=SUMPRODUCT(('Per task'!{corr_col['low']}2:{corr_col['low']}{NT}=1)*('Per task'!{corr_col['max']}2:{corr_col['max']}{NT}=0))")
sm.cell(row=re_ + 3, column=1, value="Solved at every level (tasks)")
sm.cell(row=re_ + 3, column=2, value="=SUMPRODUCT(" + "*".join(f"('Per task'!{corr_col[l]}2:{corr_col[l]}{NT}=1)" for l in L) + ")")
sm.cell(row=re_ + 4, column=1, value="Tasks")
sm.cell(row=re_ + 4, column=2, value=f"=COUNTA('Per task'!A2:A{NT})")
sm.cell(row=re_ + 5, column=1, value="Clean tasks (not flagged)")
sm.cell(row=re_ + 5, column=2, value=f"=COUNTIF('Per task'!E2:E{NT},\"no\")")
sm.cell(row=re_ + 7, column=1, value=("Reading it: all five levels sit inside each other's 95% intervals, so accuracy is a statistical tie. "
                                       "Tokens and cost are not."))
sm.cell(row=re_ + 7, column=1).font = NOTE

# ---------- By group ----------
bg = wb.create_sheet("By group", 1)
groups = ["GSM8K", "MATH", "Competition", "HumanEval", "MBPP"]
header(bg, 1, ["Group", "Tasks"] + [f"Accuracy {l}" for l in L] + [f"Mean tokens {l}" for l in L], [16, 8] + [12] * 10)
for i, g in enumerate(groups, 2):
    bg.cell(row=i, column=1, value=g).font = B
    bg.cell(row=i, column=2, value=f"=COUNTIF('Per task'!$C$2:$C${NT},A{i})")
    for k, l in enumerate(L):
        c1 = bg.cell(row=i, column=3 + k, value=f"=AVERAGEIFS('Per task'!{corr_col[l]}2:{corr_col[l]}{NT},'Per task'!$C$2:$C${NT},$A{i})")
        c1.number_format = "0.0%"
        c2 = bg.cell(row=i, column=8 + k, value=f"=AVERAGEIFS('Per task'!{tok_col[l]}2:{tok_col[l]}{NT},'Per task'!$C$2:$C${NT},$A{i})")
        c2.number_format = "#,##0"
bg.cell(row=8, column=1, value="Competition = AMC 2023, AIME 2024, OlympiadBench. MATH = levels 1 to 5.").font = NOTE

# ---------- Policies ----------
po = wb.create_sheet("Policies", 2)
header(po, 1, ["Policy", "Accuracy", "Tokens per task", "Calls per task", "Tokens vs max", "What it needs"], [40, 12, 14, 12, 12, 46])
need = {"fixed": "Nothing", "oracle": "The answer key (impossible in practice, upper bound)", "cascade": "A check you can run (e.g. visible test)",
        "length": "Nothing: uses Beam's own thinking length", "router": "BeamDial router weights", "vote": "Extra samples"}
rows = [p for p in summary["policies"] if p["kind"] != "router" or p.get("target") in (0.8, 0.9)]
mx = next(p for p in summary["policies"] if p["policy"] == "fixed: max")["tokens_per_task"]
for i, p in enumerate(rows, 2):
    po.cell(row=i, column=1, value=p["policy"]).font = N
    po.cell(row=i, column=2, value=p["accuracy"]).number_format = "0.0%"
    po.cell(row=i, column=3, value=round(p["tokens_per_task"])).number_format = "#,##0"
    po.cell(row=i, column=4, value=round(p["calls_per_task"], 2))
    po.cell(row=i, column=5, value=f"=C{i}/$C${2 + [r['policy'] for r in rows].index('fixed: max')}").number_format = "0%"
    po.cell(row=i, column=6, value=need[p["kind"]]).font = N
    for j in (2, 3, 4):
        po.cell(row=i, column=j).font = INP
po.cell(row=len(rows) + 3, column=1, value=("Source: results/summary.json from scripts/analyze.py. Every policy is replayed on the real outputs "
                                             "(each task was run at every level). Router rows are out-of-fold (cv5 = 5-fold, loso = unseen task type).")).font = NOTE

# ---------- Flagged ----------
fl = wb.create_sheet("Flagged tasks", 3)
header(fl, 1, ["Task", "Why it is flagged"], [18, 110])
for i, (t, why) in enumerate(flags.items(), 2):
    fl.cell(row=i, column=1, value=t).font = N
    fl.cell(row=i, column=2, value=why).font = N
fl.cell(row=len(flags) + 3, column=1, value="Kept in the data, reported separately. Accuracy on clean tasks is in Summary column E.").font = NOTE

# ---------- v0.2 cost plan ----------
cp = wb.create_sheet("v0.2 cost plan", 4)
cp.column_dimensions["A"].width = 44
for c in "BCDEFG":
    cp.column_dimensions[c].width = 15
cp["A1"] = "v0.2 cost plan (estimates; edit the yellow cells)"
cp["A1"].font = Font(name=F, bold=True, size=13)
inputs = [
    ("Output price (USD per 1M)", "='Summary'!B6", "Linked to Summary"),
    ("Input price (USD per 1M)", "='Summary'!B5", "Linked to Summary"),
    ("Mean output tokens per call, AIME 2024 (measured)", None, "Measured on v0.1, average over the 5 levels"),
    ("Frontier difficulty multiplier, low case", 1.0, "Assumption: frontier sets at least as hard as AIME 2024"),
    ("Frontier difficulty multiplier, high case", 2.5, "Assumption: HMMT/Apex can be ~2.5x longer"),
    ("Frontier tasks", 122, "AIME 2026 30 + HMMT Feb 2026 33 + Apex 12 + Apex Shortlist 47"),
    ("Long-context tasks", 60, "scripts/build_longctx.py"),
    ("Long-context input tokens per full pass (approx)", 1227096, "From the generator, chars/4"),
    ("Long-context tokenizer correction", 1.15, "Assumption: real tokens ~15% above chars/4"),
    ("Long-context output tokens per call, low case", 1500, "Assumption"),
    ("Long-context output tokens per call, high case", 5000, "Assumption"),
    ("Prompt tokens per frontier call", 400, "Assumption, from v0.1 prompts"),
    ("Spent so far (estimate)", 3.88, "Pilot 0.22 + probe 0.15 + full run 3.51 at the assumed prices"),
    ("Credit", 25, "Hugging Face x Reflection coupon"),
]
aime = [t for t in ids if tasks[t]["source"] == "aime24"]
aime_mean = sum(D[t][l]["usage"]["completion_tokens"] for t in aime for l in L) / (len(aime) * 5)
for i, (lab, val, note) in enumerate(inputs, 3):
    cp.cell(row=i, column=1, value=lab).font = N
    v = round(aime_mean) if val is None else val
    c = cp.cell(row=i, column=2, value=v)
    if isinstance(v, str):
        c.font = Font(name=F, color="008000")
    else:
        c.font, c.fill = INP, YEL
    cp.cell(row=i, column=3, value=note).font = NOTE
cp["B3"].number_format = cp["B4"].number_format = "$0.00"
cp["B15"].number_format = cp["B16"].number_format = "$0.00"
cp["B10"].number_format = "#,##0"
# plan table
pr = 19
header(cp, pr, ["Plan", "Samples, all levels", "Extra low samples", "Long-context samples", "Calls", "Cost, low case", "Cost, high case"])
plans = [("A. Lean", 1, 0, 1), ("B. Solid (recommended)", 1, 2, 1), ("C. Strong", 2, 2, 2)]
for i, (name, sa, el, ls) in enumerate(plans, pr + 1):
    cp.cell(row=i, column=1, value=name).font = B
    for j, v in ((2, sa), (3, el), (4, ls)):
        c = cp.cell(row=i, column=j, value=v)
        c.font, c.fill = INP, YEL
    cp.cell(row=i, column=5, value=f"=$B$8*(5*B{i}+C{i})+$B$9*5*D{i}")
    # frontier output: tasks*(5*samples*mean + extra_low*mean_low)*mult ; approximate low-level mean by measured mean
    for col, mult in ((6, "$B$6"), (7, "$B$7")):
        out_tok = "$B$10" if col == 6 else "$B$13"
        fr = f"$B$8*(5*B{i}+C{i})*$B$5*{mult}*$B$3/1000000 + $B$8*(5*B{i}+C{i})*$B$14*$B$4/1000000"
        lc = f"D{i}*5*$B$10*$B$11*$B$4/1000000 + D{i}*5*$B$9*{('$B$12' if col == 6 else '$B$13')}*$B$3/1000000"
        cp.cell(row=i, column=col, value=f"={fr} + {lc}").number_format = "$0.00"
    cp.cell(row=i, column=5).number_format = "#,##0"
cp.cell(row=pr + 5, column=1, value="Credit left after v0.1 (estimate)").font = B
cp.cell(row=pr + 5, column=2, value="=B16-B15").number_format = "$0.00"
cp.cell(row=pr + 6, column=1, value=("Extra low samples are costed at the all-level mean, so B and C are slightly conservative. "
                                     "The collector's --budget-usd is a hard stop either way.")).font = NOTE

# ---------- About ----------
ab = wb.create_sheet("About", 5)
ab.column_dimensions["A"].width = 120
lines = [
    "BeamDial results workbook",
    "",
    "Sheets: Summary (headline per level), By group, Policies (strategies replayed on real outputs), Flagged tasks, v0.2 cost plan, Per task (raw).",
    "Blue text = data or inputs. Yellow cells = assumptions you can edit. Black = formulas. Green = link to another sheet.",
    "Data: 410 tasks (GSM8K 40, MATH L1-5 160, AMC 2023 40, AIME 2024 30, OlympiadBench 40, HumanEval 50, MBPP 50), each run at all 5 reasoning levels.",
    "Grading: math with math-verify on the last \\boxed{} answer, code with the benchmark's own tests in a sandbox.",
    "Grader fix: an earlier grading under-counted math (93.1%). This workbook uses the fixed grader (96.7% over all runs).",
    "Code: https://github.com/CoderRahul01/beamdial  Data: https://huggingface.co/datasets/PengoSword123/beamdial-v0.1",
]
for i, t in enumerate(lines, 1):
    ab.cell(row=i, column=1, value=t).font = Font(name=F, bold=(i == 1), size=13 if i == 1 else 10)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title in ("Per task",)
out = ROOT / "results" / "BeamDial-results-v0.1.xlsx"
wb.save(out)
print("wrote", out)
